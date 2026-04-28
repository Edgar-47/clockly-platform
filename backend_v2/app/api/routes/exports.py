from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Literal
from uuid import UUID

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from sqlalchemy.orm import Session

from app.core.errors import ConflictError
from app.core.timezones import to_tenant_timezone
from app.db.session import get_db
from app.dependencies.auth import TenantContext, require_permission
from app.models.attendance_session import AttendanceSession
from app.models.enums import AttendanceStatus, ClockOutSource
from app.services.export_service import ExportService
from app.services.plans import check_plan_feature
from app.services.salary_service import SalaryService


router = APIRouter(prefix="/exports", tags=["exports"])


HEADERS = [
    "Empleado",
    "DNI / ID",
    "Fecha entrada",
    "Hora entrada",
    "Fecha salida",
    "Hora salida",
    "Duracion (hh:mm)",
    "Estado",
    "Origen salida",
    "Incidencia",
    "Notas",
]


@router.get("/attendance")
def export_attendance(
    export_format: Literal["excel", "xlsx", "pdf"] = Query(default="xlsx", alias="format"),
    employee_id: UUID | None = Query(default=None),
    status: AttendanceStatus | None = Query(default=AttendanceStatus.CLOSED),
    clock_out_source: ClockOutSource | None = Query(default=None),
    date_from: datetime | None = Query(default=None),
    date_to: datetime | None = Query(default=None),
    ctx: TenantContext = Depends(require_permission("exports:read")),
    db: Session = Depends(get_db),
) -> Response:
    check_plan_feature(db, ctx.company_id, "has_exports", actor_user_id=ctx.user.id)
    if employee_id is not None or date_from is not None or date_to is not None:
        check_plan_feature(db, ctx.company_id, "has_advanced_filters", actor_user_id=ctx.user.id)

    sessions = ExportService(db, company_id=ctx.company_id).list_exportable_sessions(
        employee_id=employee_id,
        status=status,
        clock_out_source=clock_out_source,
        date_from=date_from,
        date_to=date_to,
    )
    rows = [_attendance_row(session, ctx.company.timezone) for session in sessions]
    filename_base = f"clockly-fichajes-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}"
    range_label = _range_label(date_from, date_to, ctx.company.timezone)

    if export_format == "pdf":
        content = _build_pdf(
            company_name=ctx.company.name,
            timezone=ctx.company.timezone,
            range_label=range_label,
            rows=rows,
        )
        media_type = "application/pdf"
        filename = f"{filename_base}.pdf"
    else:
        content = _build_xlsx(
            company_name=ctx.company.name,
            timezone=ctx.company.timezone,
            range_label=range_label,
            rows=rows,
        )
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"{filename_base}.xlsx"

    db.commit()
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _attendance_row(session: AttendanceSession, timezone: str) -> dict[str, str]:
    clock_in = to_tenant_timezone(session.clock_in, timezone)
    clock_out = to_tenant_timezone(session.clock_out, timezone)
    return {
        "Empleado": session.employee.full_name if session.employee else str(session.employee_id),
        "DNI / ID": session.employee.dni if session.employee and session.employee.dni else str(session.employee_id),
        "Fecha entrada": clock_in.strftime("%Y-%m-%d") if clock_in else "",
        "Hora entrada": clock_in.strftime("%H:%M") if clock_in else "",
        "Fecha salida": clock_out.strftime("%Y-%m-%d") if clock_out else "",
        "Hora salida": clock_out.strftime("%H:%M") if clock_out else "",
        "Duracion (hh:mm)": _format_duration(session.duration_seconds),
        "Estado": session.status.value,
        "Origen salida": session.clock_out_source.value if session.clock_out_source else "",
        "Incidencia": "Desfichaje automatico" if session.incident_type else "",
        "Notas": session.notes or "",
    }


@router.get("/salary-calculation")
def export_salary_calculation(
    export_format: Literal["excel", "xlsx", "pdf"] = Query(default="xlsx", alias="format"),
    employee_id: UUID = Query(...),
    period_start: date = Query(..., alias="from"),
    period_end: date = Query(..., alias="to"),
    ctx: TenantContext = Depends(require_permission("salary:read")),
    db: Session = Depends(get_db),
) -> Response:
    check_plan_feature(db, ctx.company_id, "has_exports", actor_user_id=ctx.user.id)
    calculation = SalaryService(db, company_id=ctx.company_id).calculate(
        employee_id=employee_id,
        period_start=period_start,
        period_end=period_end,
    )
    filename_base = f"clockly-salarios-{datetime.now(UTC).strftime('%Y%m%d-%H%M%S')}"
    if export_format == "pdf":
        content = _build_salary_pdf(company_name=ctx.company.name, calculation=calculation)
        media_type = "application/pdf"
        filename = f"{filename_base}.pdf"
    else:
        content = _build_salary_xlsx(company_name=ctx.company.name, calculation=calculation)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        filename = f"{filename_base}.xlsx"
    return Response(
        content=content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _range_label(date_from: datetime | None, date_to: datetime | None, timezone: str) -> str:
    start = to_tenant_timezone(date_from, timezone) if date_from else None
    end = to_tenant_timezone(date_to, timezone) if date_to else None
    if start and end:
        return f"{start.strftime('%Y-%m-%d')} a {end.strftime('%Y-%m-%d')}"
    if start:
        return f"Desde {start.strftime('%Y-%m-%d')}"
    if end:
        return f"Hasta {end.strftime('%Y-%m-%d')}"
    return "Todos los registros"


def _format_duration(total_seconds: int | None) -> str:
    total_seconds = total_seconds or 0
    hours, remainder = divmod(total_seconds, 3600)
    minutes = remainder // 60
    return f"{hours:02d}:{minutes:02d}"


def _build_xlsx(*, company_name: str, timezone: str, range_label: str, rows: list[dict[str, str]]) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ConflictError("XLSX export requires the openpyxl package.") from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fichajes"

    sheet["A1"] = "ClockLy"
    sheet["A1"].font = Font(bold=True, size=16, color="0F172A")
    sheet["A2"] = company_name
    sheet["A3"] = f"Rango: {range_label}"
    sheet["A4"] = f"Zona horaria: {timezone}"
    last_column = get_column_letter(len(HEADERS))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet.merge_cells(f"A3:{last_column}3")
    sheet.merge_cells(f"A4:{last_column}4")

    header_row = 6
    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for row_index, row in enumerate(rows, start=header_row + 1):
        for column, header in enumerate(HEADERS, start=1):
            cell = sheet.cell(row=row_index, column=column, value=row[header])
            cell.alignment = Alignment(vertical="top", wrap_text=header == "Notas")
            if "Fecha" in header:
                cell.number_format = "yyyy-mm-dd"
            if "Hora" in header or header == "Duracion (hh:mm)":
                cell.alignment = Alignment(horizontal="center")

    widths = [28, 18, 15, 13, 15, 13, 17, 14, 16, 22, 42]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A{header_row}:{last_column}{max(header_row, header_row + len(rows))}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_pdf(*, company_name: str, timezone: str, range_label: str, rows: list[dict[str, str]]) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ModuleNotFoundError as exc:
        raise ConflictError("PDF export requires the reportlab package.") from exc

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
        title=f"Fichajes - {company_name}",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>ClockLy</b>", styles["Title"]),
        Paragraph(f"<b>{company_name}</b>", styles["Heading2"]),
        Paragraph(f"Informe de fichajes | Rango: {range_label} | Zona horaria: {timezone}", styles["Normal"]),
        Spacer(1, 14),
    ]

    table_rows: list[list[str]] = [HEADERS]
    for row in rows:
        table_rows.append([row[header] for header in HEADERS])
    if len(table_rows) == 1:
        table_rows.append(["Sin registros", "", "", "", "", "", "", "", ""])

    table = Table(table_rows, repeatRows=1, colWidths=[105, 78, 68, 52, 68, 52, 70, 56, 62, 88, 140])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (2, 1), (7, -1), "CENTER"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return output.getvalue()


def _build_salary_xlsx(*, company_name: str, calculation) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ModuleNotFoundError as exc:
        raise ConflictError("XLSX export requires the openpyxl package.") from exc

    headers = [
        "Empleado ID",
        "Periodo",
        "Tipo",
        "Tarifa",
        "Horas",
        "Dias",
        "Turnos",
        "Incidencias",
        "Total estimado",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Salarios estimados"
    sheet["A1"] = "ClockLy - Salarios estimados"
    sheet["A1"].font = Font(bold=True, size=16, color="0F172A")
    sheet["A2"] = company_name
    sheet["A3"] = calculation.warning
    last_column = get_column_letter(len(headers))
    sheet.merge_cells(f"A1:{last_column}1")
    sheet.merge_cells(f"A2:{last_column}2")
    sheet.merge_cells(f"A3:{last_column}3")

    header_row = 5
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for row_index, line in enumerate(calculation.lines, start=header_row + 1):
        values = [
            str(calculation.employee_id),
            f"{line.period_start} a {line.period_end}",
            line.salary_type.value,
            f"{line.amount} {line.currency}",
            str(line.total_hours),
            line.total_days,
            line.total_shifts,
            line.incident_count,
            f"{line.gross_estimated_amount} {line.currency}",
        ]
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    summary_row = header_row + len(calculation.lines) + 2
    sheet.cell(row=summary_row, column=1, value="Total estimado").font = Font(bold=True)
    sheet.cell(row=summary_row, column=9, value=f"{calculation.gross_estimated_amount} {calculation.currency}").font = Font(bold=True)
    for index, width in enumerate([38, 24, 16, 16, 12, 10, 10, 14, 18], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A6"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _build_salary_pdf(*, company_name: str, calculation) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ModuleNotFoundError as exc:
        raise ConflictError("PDF export requires the reportlab package.") from exc

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24,
        title=f"Salarios estimados - {company_name}",
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("<b>ClockLy - Salarios estimados</b>", styles["Title"]),
        Paragraph(f"<b>{company_name}</b>", styles["Heading2"]),
        Paragraph(
            f"Periodo: {calculation.period_start} a {calculation.period_end} | {calculation.warning}",
            styles["Normal"],
        ),
        Spacer(1, 14),
    ]
    headers = ["Tipo", "Tarifa", "Horas", "Dias", "Turnos", "Inc.", "Total estimado"]
    table_rows = [headers]
    for line in calculation.lines:
        table_rows.append(
            [
                line.salary_type.value,
                f"{line.amount} {line.currency}",
                str(line.total_hours),
                str(line.total_days),
                str(line.total_shifts),
                str(line.incident_count),
                f"{line.gross_estimated_amount} {line.currency}",
            ]
        )
    table_rows.append(["", "", "", "", "", "Total", f"{calculation.gross_estimated_amount} {calculation.currency}"])
    table = Table(table_rows, repeatRows=1, colWidths=[85, 85, 60, 50, 55, 45, 105])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (5, -1), (-1, -1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F8FAFC")]),
                ("ALIGN", (2, 1), (-1, -1), "CENTER"),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return output.getvalue()
