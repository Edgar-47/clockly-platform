from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Sequence

from app.core.errors import ConflictError


NAVY = "0F172A"
BLUE = "2563EB"
BLUE_DARK = "1D4ED8"
SLATE = "475569"
SLATE_LIGHT = "E2E8F0"
SURFACE = "F8FAFC"
WHITE = "FFFFFF"

DATE_FORMAT = "yyyy-mm-dd"
TIME_FORMAT = "hh:mm"
MONEY_FORMAT = '#,##0.00'
NUMBER_FORMAT = '#,##0.00'


@dataclass(frozen=True)
class XlsxKit:
    Workbook: Any
    Alignment: Any
    Border: Any
    Font: Any
    PatternFill: Any
    Side: Any
    Table: Any
    TableStyleInfo: Any
    get_column_letter: Any


def require_xlsx_kit(error_message: str) -> XlsxKit:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ModuleNotFoundError as exc:
        raise ConflictError(error_message) from exc

    return XlsxKit(
        Workbook=Workbook,
        Alignment=Alignment,
        Border=Border,
        Font=Font,
        PatternFill=PatternFill,
        Side=Side,
        Table=Table,
        TableStyleInfo=TableStyleInfo,
        get_column_letter=get_column_letter,
    )


def prepare_worksheet(sheet: Any) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def write_report_title(
    sheet: Any,
    kit: XlsxKit,
    *,
    title: str,
    company_name: str,
    details: Sequence[str],
    last_column: int,
) -> int:
    last_letter = kit.get_column_letter(last_column)
    generated_at = datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")

    sheet.merge_cells(f"A1:{last_letter}1")
    sheet.merge_cells(f"A2:{last_letter}2")
    sheet.merge_cells(f"A3:{last_letter}3")

    sheet["A1"] = title
    sheet["A1"].fill = kit.PatternFill("solid", fgColor=NAVY)
    sheet["A1"].font = kit.Font(bold=True, size=16, color=WHITE)
    sheet["A1"].alignment = kit.Alignment(vertical="center")

    sheet["A2"] = company_name
    sheet["A2"].font = kit.Font(bold=True, size=12, color=NAVY)
    sheet["A2"].alignment = kit.Alignment(vertical="center")

    detail_text = " | ".join([*details, f"Generado: {generated_at}"])
    sheet["A3"] = detail_text
    sheet["A3"].font = kit.Font(size=10, color=SLATE)
    sheet["A3"].alignment = kit.Alignment(vertical="center", wrap_text=True)

    sheet.row_dimensions[1].height = 28
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[3].height = 32
    return 5


def write_kpi_cards(
    sheet: Any,
    kit: XlsxKit,
    *,
    start_row: int,
    metrics: Sequence[tuple[str, str]],
    last_column: int,
) -> int:
    if not metrics:
        return start_row

    thin = kit.Side(style="thin", color=SLATE_LIGHT)
    border = kit.Border(left=thin, right=thin, top=thin, bottom=thin)
    card_width = max(2, last_column // len(metrics))

    for index, (label, value) in enumerate(metrics):
        start_col = index * card_width + 1
        if start_col > last_column:
            break
        end_col = last_column if index == len(metrics) - 1 else min(last_column, start_col + card_width - 1)
        start_letter = kit.get_column_letter(start_col)
        end_letter = kit.get_column_letter(end_col)
        sheet.merge_cells(f"{start_letter}{start_row}:{end_letter}{start_row}")
        sheet.merge_cells(f"{start_letter}{start_row + 1}:{end_letter}{start_row + 1}")

        for row in range(start_row, start_row + 2):
            for column in range(start_col, end_col + 1):
                cell = sheet.cell(row=row, column=column)
                cell.fill = kit.PatternFill("solid", fgColor=SURFACE)
                cell.border = border

        label_cell = sheet.cell(row=start_row, column=start_col, value=label)
        label_cell.font = kit.Font(bold=True, size=9, color=SLATE)
        label_cell.alignment = kit.Alignment(horizontal="center", vertical="center", wrap_text=True)

        value_cell = sheet.cell(row=start_row + 1, column=start_col, value=value)
        value_cell.font = kit.Font(bold=True, size=14, color=NAVY)
        value_cell.alignment = kit.Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet.row_dimensions[start_row].height = 22
    sheet.row_dimensions[start_row + 1].height = 28
    return start_row + 3


def write_table_header(sheet: Any, kit: XlsxKit, *, header_row: int, headers: Sequence[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=header)
        cell.font = kit.Font(bold=True, color=WHITE)
        cell.fill = kit.PatternFill("solid", fgColor=BLUE)
        cell.alignment = kit.Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 28


def add_excel_table(
    sheet: Any,
    kit: XlsxKit,
    *,
    name: str,
    header_row: int,
    last_row: int,
    last_column: int,
) -> None:
    last_letter = kit.get_column_letter(last_column)
    table = kit.Table(displayName=name, ref=f"A{header_row}:{last_letter}{last_row}")
    table.tableStyleInfo = kit.TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def style_table_body(
    sheet: Any,
    kit: XlsxKit,
    *,
    first_row: int,
    last_row: int,
    last_column: int,
    center_columns: set[int] | None = None,
    wrap_columns: set[int] | None = None,
) -> None:
    center_columns = center_columns or set()
    wrap_columns = wrap_columns or set()
    thin = kit.Side(style="thin", color=SLATE_LIGHT)
    border = kit.Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in sheet.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=last_column):
        for cell in row:
            cell.border = border
            cell.alignment = kit.Alignment(
                horizontal="center" if cell.column in center_columns else "left",
                vertical="top",
                wrap_text=cell.column in wrap_columns,
            )


def set_column_widths(sheet: Any, kit: XlsxKit, widths: Sequence[float]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[kit.get_column_letter(index)].width = width


def set_number_format(sheet: Any, *, column: int, first_row: int, last_row: int, number_format: str) -> None:
    for row in range(first_row, last_row + 1):
        sheet.cell(row=row, column=column).number_format = number_format


def apply_status_style(cell: Any, kit: XlsxKit) -> None:
    styles = {
        "abierto": ("DBEAFE", "1D4ED8"),
        "open": ("DBEAFE", "1D4ED8"),
        "cerrado": ("DCFCE7", "166534"),
        "closed": ("DCFCE7", "166534"),
        "anulado": ("F1F5F9", "475569"),
        "void": ("F1F5F9", "475569"),
        "pendiente": ("FEF3C7", "92400E"),
        "pending": ("FEF3C7", "92400E"),
        "en revision": ("DBEAFE", "1D4ED8"),
        "en revisión": ("DBEAFE", "1D4ED8"),
        "in_review": ("DBEAFE", "1D4ED8"),
        "aprobado": ("DCFCE7", "166534"),
        "approved": ("DCFCE7", "166534"),
        "rechazado": ("FEE2E2", "991B1B"),
        "rejected": ("FEE2E2", "991B1B"),
        "pagado": ("E0E7FF", "3730A3"),
        "paid": ("E0E7FF", "3730A3"),
    }
    key = str(cell.value or "").strip().lower()
    fill, color = styles.get(key, ("F1F5F9", SLATE))
    cell.fill = kit.PatternFill("solid", fgColor=fill)
    cell.font = kit.Font(bold=True, color=color)
    cell.alignment = kit.Alignment(horizontal="center", vertical="center", wrap_text=True)


def metric_number(value: float, *, decimals: int = 2) -> str:
    formatted = f"{value:,.{decimals}f}"
    return formatted.replace(",", "X").replace(".", ",").replace("X", ".")
